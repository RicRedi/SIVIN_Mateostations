"""Tests of the legacy workbook parser on the SYNTHETIC fixture ``legacy/data.xlsx``."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from sivin.core.ids import SensorId
from sivin.ingest.parsers.columns import ParserSettings
from sivin.ingest.parsers.legacy import LegacyWorkbookParser
from sivin.ingest.validation import Severity

from ..conftest import EXPORTS

WORKBOOK = EXPORTS / "legacy" / "data.xlsx"


def test_two_sheets_two_series(settings: ParserSettings) -> None:
    result = LegacyWorkbookParser(settings).parse(WORKBOOK)
    assert result.is_accepted, result.report.summary()
    assert result.report.issues == ()
    first, second = result.series
    assert (first.sensor_id, second.sensor_id) == (SensorId("77678271"), SensorId("77680921"))
    assert (len(first), len(second)) == (24, 24)
    sheet = openpyxl.load_workbook(WORKBOOK)["0921"]
    time_text, temp_text, rh_text = (cell.value for cell in sheet[2])
    assert time_text == "1.1.2026 01:10:42"  # CET, i.e. 00:10:42 UTC
    frame = second.frame
    assert frame["timestamp_utc"].iloc[0] == pd.Timestamp("2026-01-01 00:10:42", tz="UTC")
    assert frame["temp_c"].iloc[0] == pytest.approx(float(temp_text.replace(",", ".")))
    assert frame["rh_pct"].iloc[0] == pytest.approx(float(rh_text.replace(",", ".")))
    assert first.frame["timestamp_utc"].iloc[0] == pd.Timestamp("2026-01-01 00:00:11", tz="UTC")


def test_explicit_mapping_with_missing_and_unmapped_sheets() -> None:
    mapping = {"8271": SensorId("77678271"), "9986": SensorId("77799986")}
    parser = LegacyWorkbookParser(sheet_sensors=mapping)
    assert parser.sheet_sensors == mapping
    result = parser.parse(WORKBOOK)
    assert result.series == ()
    errors = [issue.message for issue in result.report.errors]
    warnings = [issue.message for issue in result.report.warnings]
    assert errors == ["Expected table(s) not found: 9986."]
    assert warnings == ["Table(s) without a sensor mapping were not read: 0921."]


def test_workbook_without_mapped_sheet(tmp_path: Path) -> None:
    path = tmp_path / "data.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "Notes"
    workbook.save(path)
    parser = LegacyWorkbookParser(sheet_sensors={"8271": SensorId("77678271")})
    report = parser.parse(path).report
    assert report.rules(Severity.ERROR) == {"expected-tables"}


def test_no_tables_at_all(tmp_path: Path) -> None:
    path = tmp_path / "data.xlsx"
    openpyxl.Workbook().save(path)
    parser = LegacyWorkbookParser(sheet_sensors={})
    report = parser.parse(path).report
    assert [issue.message for issue in report.errors] == ["The file contains no table to read."]


def test_can_parse(settings: ParserSettings) -> None:
    parser = LegacyWorkbookParser(settings)
    assert parser.can_parse(Path("data.xlsx"))
    assert not parser.can_parse(Path("MeteoData_8615620 77678271.xlsx"))
    assert not parser.can_parse(Path("data.csv"))
    assert not LegacyWorkbookParser().can_parse(Path("data.xlsx"))


def test_corrupt_workbook(tmp_path: Path, settings: ParserSettings) -> None:
    path = tmp_path / "data.xlsx"
    path.write_bytes(b"PK\x03\x04 not really a zip archive")
    assert LegacyWorkbookParser(settings).parse(path).report.rules() == {"file-readable"}

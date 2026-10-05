"""Edge cases of the raw cell readers and the shared reader (synthetic files)."""

from __future__ import annotations

import zipfile
from pathlib import Path

import openpyxl
import pytest

from sivin.ingest.parsers.columns import ParserSettings
from sivin.ingest.parsers.portal import PortalCsvParser, PortalXlsxParser
from sivin.ingest.parsers.sources import CsvGridReader, GridReadError, WorkbookReader
from sivin.ingest.parsers.tabular import TabularExportReader
from sivin.ingest.validation import TableInspection

from ..conftest import PORTAL_CSV_NAME

XLSX_NAME = "MeteoData_8615620 77678271.xlsx"


def test_csv_reader_reports_unreadable_file(tmp_path: Path) -> None:
    with pytest.raises(GridReadError):
        CsvGridReader(";", ("utf-8",)).read(tmp_path / "missing.csv")


def test_csv_reader_cells(tmp_path: Path) -> None:
    path = tmp_path / "a.csv"
    path.write_bytes(b'Meteo Data;\r\n"a;b";c\r\n')
    grid = CsvGridReader(";", ("utf-8",)).read(path)
    assert grid.name == "a.csv"
    assert grid.rows == (("Meteo Data", ""), ("a;b", "c"))


def test_workbook_rows_keep_their_numbers(tmp_path: Path) -> None:
    path = tmp_path / XLSX_NAME
    workbook = openpyxl.Workbook()
    workbook.active["A3"] = "Datum a čas"
    workbook.save(path)
    (grid,) = WorkbookReader().read(path)
    assert grid.rows[2][0] == "Datum a čas"
    assert len(grid.rows) == 3


def test_damaged_worksheet_xml(tmp_path: Path) -> None:
    valid = tmp_path / "valid.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["Datum a čas", "Teplota", "Vlhkost"])
    workbook.save(valid)
    damaged = tmp_path / XLSX_NAME
    with zipfile.ZipFile(valid) as source, zipfile.ZipFile(damaged, "w") as target:
        for item in source.infolist():
            data = source.read(item)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data[: len(data) // 2]
            target.writestr(item, data)
    with pytest.raises(GridReadError, match="damaged worksheet"):
        WorkbookReader().read(damaged)
    assert PortalXlsxParser().parse(damaged).report.rules() == {"file-readable"}


def test_workbook_with_chartsheet_only(tmp_path: Path) -> None:
    path = tmp_path / XLSX_NAME
    workbook = openpyxl.Workbook()
    workbook.create_chartsheet("Chart")
    workbook.remove(workbook["Sheet"])
    workbook.save(path)
    report = PortalXlsxParser().parse(path).report
    assert report.rules() == {"file-readable"}  # openpyxl cannot read a sheetless workbook


def test_unusable_path(tmp_path: Path) -> None:
    result = PortalCsvParser().parse(tmp_path / ("x" * 300) / PORTAL_CSV_NAME)
    assert result.report.rules() == {"file-exists"}


def test_assemble_requires_a_validated_table() -> None:
    reader = TabularExportReader(ParserSettings())
    with pytest.raises(ValueError, match="not validated"):
        reader.assemble(TableInspection(name="t", sensor_id=None), "a.csv")
